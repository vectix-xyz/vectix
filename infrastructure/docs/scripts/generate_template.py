from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def create_docx_template():
    doc = Document()

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    cell_left = header_table.cell(0, 0)
    p_logo = cell_left.paragraphs[0]
    run_logo = p_logo.add_run("BusToMove")
    run_logo.font.name = "Arial"
    run_logo.font.size = Pt(18)
    run_logo.font.bold = True
    run_logo.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_sub = cell_left.add_paragraph()
    run_sub = p_sub.add_run("Document Generation Service • Automatically generated")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(8.5)
    run_sub.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    cell_right = header_table.cell(0, 1)
    p_date = cell_right.paragraphs[0]
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_date_lbl = p_date.add_run("Report date:\n")
    run_date_lbl.font.size = Pt(8)
    run_date_lbl.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    run_date = p_date.add_run("{{ report_date }}")
    run_date.font.size = Pt(10)
    run_date.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    title_table = doc.add_table(rows=1, cols=2)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    p_title = title_table.cell(0, 0).paragraphs[0]
    run_title = p_title.add_run("Trip report for: {{ company_name }}")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_badge = title_table.cell(0, 1).paragraphs[0]
    p_badge.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_badge = p_badge.add_run("Total trips: {{ trips|length }}")
    run_badge.font.name = "Arial"
    run_badge.font.size = Pt(9)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Час", "Маршрут", "Транспорт", "Пасажири"]
    widths = [Inches(1.0), Inches(2.8), Inches(1.8), Inches(1.0)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "1E3A8A")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i == 3 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.font.name = "Arial"
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    row_cells = table.rows[1].cells

    p0 = row_cells[0].paragraphs[0]
    p0.add_run("{% tr for trip in trips %}\n{{ trip.time }}")

    p1 = row_cells[1].paragraphs[0]
    p1.add_run("{{ trip.route }}")

    p2 = row_cells[2].paragraphs[0]
    p2.add_run("{{ trip.bus }}")

    p3 = row_cells[3].paragraphs[0]
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p3.add_run("{{ trip.passengers }}\n{% tr endfor %}")

    doc.add_paragraph().paragraph_format.space_before = Pt(36)

    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    p_sig1 = sig_table.cell(0, 0).paragraphs[0]
    p_sig1.add_run("Responsible:\n\n_______________________\n(Sign / Name Surname)").font.size = Pt(8.5)

    p_sig2 = sig_table.cell(0, 1).paragraphs[0]
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig2.add_run("Taken by:\n\n_______________________\n(Sign / Name Surname)").font.size = Pt(8.5)

    output_path = Path("src/infrastructure/adapters/templates/trip_schedule.docx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"✅ Шаблон успішно збережено в: {output_path.resolve()}")

if __name__ == "__main__":
    create_docx_template()